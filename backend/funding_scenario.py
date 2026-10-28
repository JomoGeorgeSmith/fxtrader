"""Opt-in, additive synthetic funding scenario. Never seed this on startup."""
import argparse
import json
import shutil
from copy import copy
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .seed import ROOT, SHARED

PREFIX = 'SIM-FUNDING-'

def scenario_rows(today):
    def row(id,currency,amount,offset,status,description,category=None):
        item=dict(id=PREFIX+id,currency=currency,amount=amount,
                  value_date=(today+timedelta(days=offset)).isoformat(),status=status,
                  description='SYNTHETIC funding experiment: '+description)
        if category:
            item['category']=category
        return item
    payments=[
        row('P01','USD',450000,2,'Committed','Short-notice loan drawdown','Loan'),
        row('P02','USD',2200000,21,'Committed','Confirmed corporate loan request; three weeks notice','Loan'),
        row('P03','USD',350000,12,'Committed','Overseas equipment settlement','Bill'),
        row('P04','USD',650000,18,'Expected','Tentative additional loan request; approval pending','Loan'),
        row('P05','JMD',90000000,9,'Committed','Local loan drawdown','Loan'),
        row('P06','JMD',120000000,23,'Committed','Payroll and operating settlements','Bill'),
        row('P07','GBP',1100000,10,'Committed','Confirmed sterling loan drawdown','Loan'),
        row('P08','GBP',650000,24,'Committed','UK supplier obligation','Bill'),
        row('P09','CAD',400000,16,'Committed','Canadian customer settlement','Currency sale'),
    ]
    receipts=[
        row('R01','USD',200000,3,'Confirmed','Scheduled customer deposit'),
        row('R02','USD',300000,6,'Confirmed','Confirmed loan repayment'),
        row('R03','USD',800000,5,'Expected','Expected customer deposit; timing unconfirmed'),
        row('R04','USD',600000,15,'Expected','Expected foreign-currency deposit; not committed'),
        row('R05','JMD',100000000,2,'Confirmed','Scheduled corporate deposit'),
        row('R06','JMD',80000000,6,'Expected','Expected customer deposit; timing unconfirmed'),
        row('R07','GBP',150000,3,'Confirmed','Scheduled sterling deposit'),
        row('R08','GBP',400000,7,'Expected','Expected sterling deposit; confirmation outstanding'),
        row('R09','CAD',200000,4,'Confirmed','Scheduled Canadian-dollar deposit'),
    ]
    return {'Payments':payments,'Expected_Receipts':receipts}

def apply_scenario(shared=SHARED,runtime=ROOT/'runtime',today=None):
    today=today or datetime.now().date()
    path=Path(shared)/'Weekly_Disbursements.xlsx'
    wb=load_workbook(path)
    try:
        if any(str(row[0]).startswith(PREFIX) for sheet in ('Payments','Expected_Receipts') for row in wb[sheet].iter_rows(min_row=2,values_only=True)):
            raise ValueError('Scenario rows already exist. No rows were appended or overwritten.')
        folder=Path(runtime)/'experiments'/('funding-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
        folder.mkdir(parents=True)
        shutil.copy2(path,folder/path.name)
        changes=scenario_rows(today)
        for name,rows in changes.items():
            ws=wb[name]
            headers=[cell.value for cell in ws[1]]
            for item in rows:
                ws.append([item.get(header) for header in headers])
                for cell in ws[ws.max_row]:
                    template=ws.cell(2,cell.column)
                    cell._style=copy(template._style)
            for table in ws.tables.values():
                table.ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
                if table.autoFilter:
                    table.autoFilter.ref=table.ref
        if 'Read_Me' in wb.sheetnames:
            wb['Read_Me'].append([f'Funding experiment added {today}: SIM-FUNDING rows only; original rows and balances unchanged.'])
            wb['Read_Me'].append(['Four-week sample only, not a complete obligation schedule. Expected receipts are uncertain; Expected payments are tentative requests. No preferred FX action is embedded.'])
        manifest=dict(start_date=today.isoformat(),backup=str(folder/path.name),added_rows=changes,
                      note='Only synthetic cash-flow rows appended. Existing rows, balances, rates and risk inputs preserved. This is not a complete cash-flow schedule.')
        (folder/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        wb.save(path)
        return manifest
    finally:
        wb.close()

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply',action='store_true',help='Append the synthetic rows to the shared workbook, after backup')
    args=parser.parse_args()
    if args.apply:
        print(json.dumps(apply_scenario(),indent=2))
    else:
        print(json.dumps(scenario_rows(datetime.now().date()),indent=2))
