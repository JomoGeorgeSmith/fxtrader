"""Generate reproducible synthetic treasury workbooks; never overwrite by default."""
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl import load_workbook
from .liquidity_metrics import COMPONENTS

ROOT = Path(__file__).resolve().parents[1]
SHARED = ROOT / 'shared'

def write_book(name, tabs, force=False):
    target = SHARED / name
    if target.exists() and not force:
        return
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = 'SYNTHETIC POC DATA — NOT ACTUAL JN BANK DATA'
    wb.properties.description = 'Demo balances, rates and limits. Not official JN or BOJ requirements.'
    for index, (title, headers, rows) in enumerate(tabs):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='124980')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(38, max(18, max(len(str(c.value or '')) for c in col) + 3))
            for cell in col[1:]:
                if isinstance(cell.value, (float, int)):
                    cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='center')
        table = Table(displayName=f'Table{index}', ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    notes = wb.create_sheet('Read_Me')
    notes.append(['SYNTHETIC POC DATA — NOT ACTUAL JN BANK DATA'])
    notes.append(['All figures, rates, accounts and risk inputs are made up for demonstration.'])
    notes.append(['Demo rules are not official JN or BOJ requirements.'])
    notes.append(['Use numeric values; formula recalculation is not performed by the backend.'])
    notes.append(['Save in Excel to trigger a new analysis. Formatting-only changes are ignored.'])
    notes.column_dimensions['A'].width = 100
    notes['A1'].font = Font(bold=True,color='124980')
    wb.save(target)

def seed():
    SHARED.mkdir(exist_ok=True)
    now = datetime.now().replace(microsecond=0)
    stamp = now.isoformat()
    write_book('Bank_Balances.xlsx', [
        ('Accounts', ['account_id','bank','currency','balance','holds','as_of'], [
            ['BOJ-JMD-OPS','BOJ','JMD',420000000,10000000,stamp],
            ['CITI-USD-OPS','Citi','USD',2400000,100000,stamp],
            ['CITI-GBP-OPS','Citi','GBP',480000,10000,stamp],
            ['CITI-CAD-OPS','Citi','CAD',650000,20000,stamp]]),
        ('Reserves', ['reserve_id','currency','balance','as_of'], [
            ['BOJ-JMD-RES','JMD',215000000,stamp], ['BOJ-USD-RES','USD',460000,stamp],
            ['BOJ-GBP-RES','GBP',60000,stamp], ['BOJ-CAD-RES','CAD',90000,stamp]]),
        ('Risk_Inputs', ['currency','net_fx_exposure','capital_base_jmd','as_of'], [
            ['USD',9000000,10000000000,stamp], ['GBP',300000,10000000000,stamp], ['CAD',350000,10000000000,stamp]])])
    write_book('Daily_Rate_Sheet.xlsx', [('Rates', ['currency','buy','sell','previous_sell','as_of','quote'], [
        ['USD',154.5,159,158.5,stamp,'JMD per USD'], ['GBP',202.18,213.94,208,stamp,'JMD per GBP'], ['CAD',107.98,113.66,113.2,stamp,'JMD per CAD']])])
    outflows = [
        ('USD',950000,1,'Loan','Corporate loan drawdown'), ('USD',1350000,3,'Currency sale','Committed customer FX sale'),
        ('USD',550000,5,'Bill','Overseas supplier settlement'), ('GBP',180000,2,'Loan','Sterling loan drawdown'),
        ('GBP',90000,4,'Bill','UK service payments'), ('CAD',150000,2,'Currency sale','Customer CAD settlement'),
        ('CAD',120000,5,'Bill','Canadian supplier payments'), ('JMD',170000000,1,'Loan','Local mortgage disbursements'),
        ('JMD',130000000,4,'Bill','Domestic operating payments')]
    receipts = [('USD',250000,2), ('GBP',70000,3), ('CAD',120000,3), ('JMD',80000000,2)]
    write_book('Weekly_Disbursements.xlsx', [
        ('Payments', ['id','currency','amount','value_date','category','description','status'], [
            [f'PAY-{i+1:03}',c,a,(now+timedelta(days=d)).date().isoformat(),cat,desc,'Committed'] for i,(c,a,d,cat,desc) in enumerate(outflows)]),
        ('Expected_Receipts', ['id','currency','amount','value_date','description','status'], [
            [f'REC-{i+1:03}',c,a,(now+timedelta(days=d)).date().isoformat(),'Scheduled receipt','Confirmed'] for i,(c,a,d) in enumerate(receipts)])])
    ensure_liquidity_inputs()

def ensure_liquidity_inputs():
    """Add one new tab to an existing demo book; preserve all existing user values."""
    path = SHARED/'Bank_Balances.xlsx'
    wb = load_workbook(path)
    try:
        if 'Liquidity_Inputs' in wb.sheetnames:
            return
        ws = wb.create_sheet('Liquidity_Inputs')
        ws.append(['component','currency','amount','as_of','description'])
        stamp = datetime.now().replace(microsecond=0).isoformat()
        for component,(amount,description) in COMPONENTS.items():
            ws.append([component,'JMD',amount,stamp,description])
        ws.freeze_panes='A2'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='124980')
        for col,width in [('A',34),('B',16),('C',24),('D',26),('E',95)]:
            ws.column_dimensions[col].width=width
        for cell in list(ws['C'])[1:]:
            cell.number_format='#,##0.00'
        table=Table(displayName='SyntheticLiquidityInputs',ref=ws.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
        wb.save(path)
    finally:
        wb.close()

if __name__ == '__main__':
    seed()
    print(f'Synthetic workbooks ready in {SHARED}')
