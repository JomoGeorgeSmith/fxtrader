"""Optional Excel extension: missing simulated FX records, not replacement balances."""
from datetime import date
from math import isfinite
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

SIMULATION_FILE='FX_Simulation.xlsx'
Currency=Literal['JMD','USD','GBP','CAD']


class Row(BaseModel):
    model_config=ConfigDict(extra='forbid',allow_inf_nan=False)
    id:str=Field(min_length=1)
    currency:Currency
    synthetic:Literal[True]
    description:str

class Limit(Row):
    max_absolute_exposure:float=Field(gt=0)

class Terms(Row):
    annual_interest_rate:float=Field(ge=0,le=1)
    spot_business_days:int=Field(ge=0,le=10)
    cutoff_hour:int=Field(ge=0,le=23)
    fee_bps:float=Field(ge=0,le=1000)
    fixed_fee_jmd:float=Field(ge=0)

class Permission(Row):
    account_id:str
    fx_enabled:bool
    daily_debit_limit:float=Field(ge=0)

class ReceiptTiming(Row):
    receipt_id:str
    confirmation_reference:str
    delay_days:int=Field(ge=0,le=28)
    arrival_hour:int=Field(ge=0,le=23)

class Holiday(Row):
    value_date:date

class Order(Row):
    source:Currency
    target:Currency
    source_amount:float=Field(gt=0)
    target_amount:float=Field(gt=0)
    settlement_date:date
    status:Literal['Planned','Booked','Settled','Cancelled']
    obligation_id:str
    cash_in_forecast:bool
    payment_id:str
    receipt_id:str
    fee_jmd:float=Field(ge=0)
    fee_payment_id:str
    cancellation_fee_jmd:float=Field(ge=0)
    amendment_fee_jmd:float=Field(ge=0)
    cancellable:bool
    amendable:bool

class HedgePolicy(Row):
    allow_deliverable_forward:bool
    max_hedge_pct:float=Field(ge=0,le=100)
    max_tenor_days:int=Field(gt=0,le=730)
    forward_spread_bps:float=Field(ge=0,le=1000)

class Counterparty(Row):
    approved:bool
    forward_limit_jmd:float=Field(gt=0)
    collateral_pct:float=Field(ge=0,le=100)
    fee_jmd:float=Field(ge=0)

class Hedge(Row):
    obligation_id:str
    source:Currency
    target:Currency
    source_amount:float=Field(gt=0)
    target_amount:float=Field(gt=0)
    settlement_date:date
    status:Literal['Active','Expired','Cancelled']
    counterparty_id:str
    collateral_jmd:float=Field(ge=0)
    collateral_release_receipt_id:str
    cash_in_forecast:bool
    payment_id:str
    receipt_id:str
    fee_jmd:float=Field(ge=0)
    fee_payment_id:str

class Allocation(Row):
    receipt_id:str
    obligation_id:str
    amount:float=Field(gt=0)


TABLES={
    'fx_limits':('FX_Limits',Limit),
    'fx_terms':('Currency_Terms',Terms),
    'fx_permissions':('Account_Permissions',Permission),
    'fx_receipt_timing':('Receipt_Timing',ReceiptTiming),
    'fx_holidays':('Settlement_Closures',Holiday),
    'fx_orders':('FX_Orders',Order),
    'fx_hedge_policy':('Hedge_Policy',HedgePolicy),
    'fx_counterparties':('Counterparties',Counterparty),
    'fx_hedges':('FX_Hedges',Hedge),
    'fx_allocations':('Receipt_Allocations',Allocation),
}


def read_simulation(folder,data):
    result={key:[] for key in TABLES}
    path=Path(folder)/SIMULATION_FILE
    if not path.exists():return result
    wb=load_workbook(path,data_only=True)
    try:
        for key,(sheet,model) in TABLES.items():
            if sheet not in wb.sheetnames:
                raise ValueError(f'{SIMULATION_FILE}: missing {sheet}')
            rows=list(wb[sheet].values)
            headers=list(rows[0])
            if set(headers)!=set(model.model_fields) or len(headers)!=len(set(headers)):
                raise ValueError(f'{sheet}: columns must match the simulation schema')
            for index,values in enumerate(rows[1:],2):
                if all(v is None for v in values):continue
                raw=dict(zip(headers,values))
                for field in model.model_fields:
                    if raw.get(field) is None and model.model_fields[field].annotation is str:raw[field]=''
                item=model.model_validate(raw).model_dump(mode='json')
                for field,value in item.items():
                    if isinstance(value,float) and not isfinite(value):raise ValueError(f'{sheet}: non-finite {field}')
                item['_source']=f'{SIMULATION_FILE} / {sheet} / row {index}'
                result[key].append(item)
        validate_links(data|result)
        return result
    finally:wb.close()


def validate_links(data):
    def by_id(key,field='id'):
        rows=data.get(key,[])
        ids=[r[field] for r in rows]
        if len(ids)!=len(set(ids)):raise ValueError(f'{key}: duplicate {field}')
        return dict(zip(ids,rows))
    accounts=by_id('accounts','account_id');payments=by_id('payments');receipts=by_id('receipts')
    for key in TABLES:by_id(key)
    ids=[r['id'] for key in TABLES for r in data.get(key,[])]
    if len(ids)!=len(set(ids)):raise ValueError('Simulation IDs must be unique across extension tabs for evidence references')
    for key in ('fx_limits','fx_terms','fx_hedge_policy'):
        by_id(key,'currency')
    if any(r['currency']=='JMD' for r in data.get('fx_limits',[])):
        raise ValueError('FX_Limits supplements the three foreign-currency Risk_Inputs positions, not JMD')
    if data.get('fx_terms') and {r['currency'] for r in data['fx_terms']}!={'JMD','USD','GBP','CAD'}:
        raise ValueError('Currency_Terms must cover all supported currencies')
    by_id('fx_permissions','account_id');by_id('fx_receipt_timing','receipt_id')
    for row in data.get('fx_permissions',[]):
        if row['account_id'] not in accounts or accounts[row['account_id']]['currency']!=row['currency']:
            raise ValueError('Account permission must reference an existing same-currency account')
    for row in data.get('fx_receipt_timing',[]):
        if row['receipt_id'] not in receipts or receipts[row['receipt_id']]['currency']!=row['currency']:
            raise ValueError('Receipt timing must reference an existing same-currency receipt')
        if row['confirmation_reference'] and receipts[row['receipt_id']]['status']=='Expected':
            raise ValueError('Expected receipts cannot have a simulated confirmation without updating their original status')
    for row in data.get('fx_allocations',[]):
        if row['receipt_id'] not in receipts or row['obligation_id'] not in payments:
            raise ValueError('Allocation refers to an unknown receipt/payment')
        if {row['currency'],receipts[row['receipt_id']]['currency'],payments[row['obligation_id']]['currency']}!={row['currency']}:
            raise ValueError('Allocation currencies must agree')
    for id,receipt in receipts.items():
        if sum((Decimal(str(r['amount'])) for r in data.get('fx_allocations',[]) if r['receipt_id']==id),Decimal(0))>Decimal(str(receipt['amount'])):
            raise ValueError('Receipt allocations exceed its amount')
    for id,payment in payments.items():
        if sum((Decimal(str(r['amount'])) for r in data.get('fx_allocations',[]) if r['obligation_id']==id),Decimal(0))>Decimal(str(payment['amount'])):
            raise ValueError('Receipt allocations exceed their payment amount')
    cps=by_id('fx_counterparties')
    used_flows=set()
    for key in ('fx_orders','fx_hedges'):
        for row in data.get(key,[]):
            if row['source']==row['target'] or row['currency']!=row['target']:
                raise ValueError('Trade source/target must differ and currency must equal target')
            if row['obligation_id'] and (row['obligation_id'] not in payments or payments[row['obligation_id']]['currency']!=row['target']):
                raise ValueError('Trade obligation must exist in the target currency')
            if key=='fx_hedges' and row['counterparty_id'] not in cps:raise ValueError('Unknown hedge counterparty')
            if key=='fx_hedges' and row['collateral_release_receipt_id']:
                release=receipts.get(row['collateral_release_receipt_id'])
                if not release or release['currency']!='JMD' or release['amount']!=row['collateral_jmd'] or release['value_date']!=row['settlement_date'] or release['status'] in ('Cancelled','Settled'):
                    raise ValueError('Collateral release must reference a matching active dated JMD receipt')
                ref=('receipt',row['collateral_release_receipt_id'])
                if ref in used_flows:raise ValueError('A collateral release cannot represent multiple cash legs')
                used_flows.add(ref)
            if row['cash_in_forecast']:
                p=payments.get(row['payment_id']);r=receipts.get(row['receipt_id'])
                if not p or not r or p['currency']!=row['source'] or r['currency']!=row['target'] or p['amount']!=row['source_amount'] or r['amount']!=row['target_amount'] or p['value_date']!=row['settlement_date'] or r['value_date']!=row['settlement_date']:
                    raise ValueError('Included trade cash legs must match existing dated flows exactly')
                if p['status'] in ('Cancelled','Settled') or r['status'] in ('Cancelled','Settled'):
                    raise ValueError('Included trade cash legs must be active')
                for ref in (('payment',row['payment_id']),('receipt',row['receipt_id'])):
                    if ref in used_flows:raise ValueError('A cash leg cannot represent multiple trades')
                    used_flows.add(ref)
            elif row['payment_id'] or row['receipt_id']:
                raise ValueError('Cash-flow references require cash_in_forecast=true')
            if row['fee_payment_id']:
                fee=payments.get(row['fee_payment_id'])
                if not fee or fee['currency']!='JMD' or fee['amount']!=row['fee_jmd'] or fee['value_date']!=row['settlement_date'] or fee['status'] in ('Cancelled','Settled'):
                    raise ValueError('Trade fee reference must match an active dated JMD payment')
                ref=('payment',row['fee_payment_id'])
                if ref in used_flows:raise ValueError('A fee cash leg cannot represent multiple trades')
                used_flows.add(ref)
            if key=='fx_orders' and row['status']=='Planned' and row['cash_in_forecast']:
                raise ValueError('Unapproved planned orders cannot already be cash commitments')
            if key=='fx_orders' and row['status']=='Planned' and row['fee_payment_id']:
                raise ValueError('Unapproved planned order fees cannot already be cash commitments')
    collateral=sum((Decimal(str(r['collateral_jmd'])) for r in data.get('fx_hedges',[]) if r['status']=='Active'),Decimal(0))
    held=sum((Decimal(str(a['holds'])) for a in data['accounts'] if a['currency']=='JMD'),Decimal(0))
    if collateral>held:raise ValueError('Existing active hedge collateral must already be included in JMD account holds')
