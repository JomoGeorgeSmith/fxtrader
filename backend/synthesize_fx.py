"""Opt-in missing-data extension. Existing balances, rates and flows are never edited."""
import argparse
import hashlib
import json
import shutil
from datetime import datetime,timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo

from .seed import ROOT,SHARED
from .simulation_inputs import SIMULATION_FILE,TABLES,read_simulation

BASE_FILES=('Bank_Balances.xlsx','Daily_Rate_Sheet.xlsx','Weekly_Disbursements.xlsx')


def digest_file(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def missing_rows(data,today):
    def row(id,currency,description,**fields):
        return dict(id='SIM-FX-'+id,currency=currency,synthetic=True,description='SYNTHETIC assumption: '+description,**fields)
    result={key:[] for key in TABLES}
    capital=Decimal(str(data['risk'][0]['capital_base_jmd']))
    rates={r['currency']:r for r in data['rates']}
    for currency in ('USD','GBP','CAD'):
        # Equal JMD-equivalent risk capacity, not a hand-picked winning trade.
        limit=(capital*Decimal('.08')/Decimal(str(rates[currency]['sell']))).quantize(Decimal('.01'))
        result['fx_limits'].append(row('LIMIT-'+currency,currency,'Per-currency exposure capacity; existing aggregate rule remains.',max_absolute_exposure=float(limit)))
    assumed_rates={'JMD':.075,'USD':.05,'GBP':.045,'CAD':.04}
    for currency,rate in assumed_rates.items():
        result['fx_terms'].append(row('TERMS-'+currency,currency,'Illustrative annual interest, processing and fees; not observed market terms.',annual_interest_rate=rate,spot_business_days=2,cutoff_hour=14,fee_bps=5,fixed_fee_jmd=1000))
        result['fx_hedge_policy'].append(row('HEDGE-POLICY-'+currency,currency,'Deliverable forwards only; no speculative or over-sized hedge.',allow_deliverable_forward=currency!='JMD',max_hedge_pct=80,max_tenor_days=365,forward_spread_bps=10))
    for a in data['accounts']:
        result['fx_permissions'].append(row('ACCOUNT-'+a['account_id'],a['currency'],'Operating account enabled; protected reserves are excluded.',account_id=a['account_id'],fx_enabled=True,daily_debit_limit=max(0,a['balance']-a['holds'])))
    for receipt in data['receipts']:
        confirmed=receipt['status'] in ('Confirmed','Committed')
        result['fx_receipt_timing'].append(row('RECEIPT-'+receipt['id'],receipt['currency'],'Timing metadata only; original amount, date and status are authoritative.',receipt_id=receipt['id'],confirmation_reference='SIM-CONF-'+receipt['id'] if confirmed else '',delay_days=1 if confirmed else 3,arrival_hour=11))
    for currency in assumed_rates:
        for year in (today.year,today.year+1):
            result['fx_holidays'].append(row(f'CLOSURE-{currency}-{year}',currency,'Assumed year-end bank closure; not an official holiday calendar.',value_date=f'{year}-12-25'))
    result['fx_counterparties'].append(row('DEALER-01','JMD','Simulated approved counterparty, not a real dealer offer.',approved=True,forward_limit_jmd=float(capital*Decimal('.15')),collateral_pct=5,fee_jmd=2500))
    # A small unapproved round trip is planning data, not an operating cash flow.
    # It offers something to assess without asserting that cancellation is best.
    foreign=[c for c in ('GBP','USD','CAD') if any(a['currency']==c for a in data['accounts'])]
    if len(foreign)>=2:
        source,target=foreign[:2]
        available=sum(Decimal(str(a['balance']))-Decimal(str(a['holds'])) for a in data['accounts'] if a['currency']==source)
        amount=min(Decimal('100000'),max(Decimal(0),available/100)).quantize(Decimal('.01'))
        if amount>0:
            received=(amount*Decimal(str(rates[source]['buy']))/Decimal(str(rates[target]['sell']))).quantize(Decimal('.01'))
            returned=(received*Decimal(str(rates[target]['buy']))/Decimal(str(rates[source]['sell']))).quantize(Decimal('.01'))
            def business(offset):
                day=today+timedelta(days=offset)
                while day.weekday()>4:day+=timedelta(days=1)
                return day.isoformat()
            for id,s,t,sold,bought,day in [('ORDER-01',source,target,amount,received,business(5)),('ORDER-02',target,source,received,returned,business(10))]:
                fee=float((sold*Decimal(str(rates[s]['buy']))*Decimal('.0005')+1000).quantize(Decimal('.01')))
                result['fx_orders'].append(row(id,t,'Unapproved planning order; amounts use existing sheet rates, not firm prices.',source=s,target=t,source_amount=float(sold),target_amount=float(bought),settlement_date=day,status='Planned',obligation_id='',cash_in_forecast=False,payment_id='',receipt_id='',fee_jmd=fee,fee_payment_id='',cancellation_fee_jmd=0,amendment_fee_jmd=0,cancellable=True,amendable=True))
    # Fill the missing planning register using existing obligations/receipts;
    # no existing cash amount, status or date is changed to make a case win.
    candidates=[p for p in data['payments'] if p['currency']!='JMD' and p['status'] not in ('Cancelled','Settled')
        and p['value_date']>=(today+timedelta(days=5)).isoformat()
        and any(r['currency']==p['currency'] and r['status'] not in ('Cancelled','Settled') and today.isoformat()<=r['value_date']<p['value_date'] for r in data['receipts'])]
    for index,p in enumerate(sorted(candidates,key=lambda p:(p['value_date'],p['id']))[:2],1):
        target=p['currency'];source='GBP' if target!='GBP' else 'USD'
        bought=Decimal(str(p['amount']));sold=(bought*Decimal(str(rates[target]['sell']))/Decimal(str(rates[source]['buy']))).quantize(Decimal('.01'),rounding='ROUND_CEILING')
        day=datetime.fromisoformat(p['value_date']).date()-timedelta(days=2)
        while day.weekday()>4:day-=timedelta(days=1)
        cost=(sold*Decimal(str(rates[source]['buy']))*Decimal('.0005')+1000).quantize(Decimal('.01'),rounding='ROUND_CEILING')
        result['fx_orders'].append(row(f'PURCHASE-{index:02}',target,'Unapproved full-payment purchase plan linked to an existing obligation; not a recommendation.',source=source,target=target,source_amount=float(sold),target_amount=float(bought),settlement_date=day.isoformat(),status='Planned',obligation_id=p['id'],cash_in_forecast=False,payment_id='',receipt_id='',fee_jmd=float(cost),fee_payment_id='',cancellation_fee_jmd=0,amendment_fee_jmd=500,cancellable=True,amendable=True))
    # Empty, schema-defined ledgers explicitly mean no recorded existing hedges
    # or earmarks in this simulation, not fabricated real-world confirmations.
    return result


def apply_missing(shared=SHARED,runtime=ROOT/'runtime',today=None):
    from .app import read_context
    shared=Path(shared);runtime=Path(runtime);today=today or datetime.now().date()
    data=read_context(shared,include_simulation=False)
    target=shared/SIMULATION_FILE
    wb=load_workbook(target) if target.exists() else Workbook()
    if not target.exists():wb.remove(wb.active)
    missing=[key for key,(sheet,_) in TABLES.items() if sheet not in wb.sheetnames]
    if not missing:
        wb.close();read_simulation(shared,data)
        return dict(unchanged=True,file=str(target),added_sheets=[])
    files=[shared/name for name in BASE_FILES]+([target] if target.exists() else [])
    hashes={p.name:digest_file(p) for p in files}
    folder=runtime/'experiments'/('fx-extension-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
    before=folder/'before';before.mkdir(parents=True)
    for p in files:shutil.copy2(p,before/p.name)
    rows=missing_rows(data,today)
    try:
        for key in missing:
            title,model=TABLES[key];ws=wb.create_sheet(title)
            headers=list(model.model_fields);ws.append(headers)
            for item in rows[key]:ws.append([item[h] for h in headers])
            ws.freeze_panes='A2'
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='124980')
            for index,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(index)].width=min(65,max(16,len(h)+3))
            if rows[key]:
                table=Table(displayName='SimFX'+title.replace('_',''),ref=ws.dimensions)
                table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
                ws.add_table(table)
        if 'Read_Me' not in wb.sheetnames:
            notes=wb.create_sheet('Read_Me')
            for note in ['SYNTHETIC FX SIMULATION — NOT ACTUAL BANK DATA',
                'Existing balances, risk inputs, reserves, rates, payments, receipts and statuses are reused unchanged.',
                'All added limits, confirmations, timings, policies, rates and orders are simulation assumptions, not real commitments.',
                'Interest rates are decimal annual rates; fees/spreads marked bps are basis points. Amounts are native currency unless explicitly JMD.',
                'Planned orders are unapproved and excluded from operating forecasts. Booked orders/Active hedges add cash legs unless exactly linked to existing flows.',
                'Settled trade cash is assumed reflected in opening balances. Risk_Inputs remains the authoritative current exposure; do not add historical trade notionals again.',
                'Empty FX_Hedges and Receipt_Allocations declare no recorded existing hedges/earmarks in this simulation.',
                'Calendar closures, processing days and cutoffs are assumed; no real delivery or executable forward price is verified.',
                'Every trade assessment is independent unless it explicitly models a combined change. No opportunity or trade is prescribed.']:
                notes.append([note])
            notes.column_dimensions['A'].width=130
        wb.properties.title='SYNTHETIC FX SIMULATION — MISSING INPUT EXTENSION'
        staged=folder/SIMULATION_FILE;wb.save(staged)
        read_simulation(folder,data)
        if any(digest_file(p)!=hashes[p.name] for p in files):raise ValueError('Workbook changed during setup; no changes applied')
        shutil.copy2(staged,target)
        manifest=dict(file=str(target),backup=str(before),added_sheets=[TABLES[k][0] for k in missing],added_rows={k:len(rows[k]) for k in missing},original_hashes=hashes,
            preserved_base_files=all(digest_file(shared/name)==hashes[name] for name in BASE_FILES),
            reused=['accounts','reserves','risk','rates','payments','receipts','liquidity_inputs','receipt statuses','aggregate/cash limits'],
            note='Only missing extension tabs added; existing tabs and original workbooks were not regenerated. No paid analysis was requested.')
        (folder/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        return manifest
    finally:wb.close()


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply',action='store_true');args=parser.parse_args()
    if args.apply:print(json.dumps(apply_missing(),indent=2))
    else:
        from .app import read_context
        print(json.dumps(missing_rows(read_context(SHARED,include_simulation=False),datetime.now().date()),indent=2))
