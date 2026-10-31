"""One-off reproducible synthetic portfolio extension; never run with the watcher active.

Only applies to the original October 27 demo snapshot. Idempotent by row ID.
The pre-change workbooks and audit live in the separate restore archive.
"""
from copy import copy
from datetime import date,datetime,timedelta
from pathlib import Path

from openpyxl import load_workbook


def extend(folder: Path, anchor: date = date(2026, 10, 30)):
    folder=Path(folder)
    files=['Bank_Balances.xlsx','Daily_Rate_Sheet.xlsx','Weekly_Disbursements.xlsx','FX_Simulation.xlsx']
    books={name:load_workbook(folder/name) for name in files}
    try:
        payments=books['Weekly_Disbursements.xlsx']['Payments']
        if any(row[0]=='POC-USD-FUTURE' for row in list(payments.values)[1:]):
            return False
        accounts=books['Bank_Balances.xlsx']['Accounts']
        headers=[c.value for c in accounts[1]]
        original=datetime.fromisoformat(str(accounts.cell(2,headers.index('as_of')+1).value)).date()
        if original!=date(2026,10,27) or anchor<original:
            raise ValueError('Not the original synthetic October 27 portfolio')
        offset=anchor-original
        for wb in books.values():
            for ws in wb:
                if ws.title=='Read_Me':continue
                columns={cell.value:cell.column for cell in ws[1]}
                for field in ('as_of','value_date','settlement_date'):
                    if field not in columns:continue
                    for row in range(2,ws.max_row+1):
                        cell=ws.cell(row,columns[field]);value=cell.value
                        if value is None or value=='':continue
                        if isinstance(value,(datetime,date)):shifted=value+offset
                        elif field=='as_of':shifted=datetime.fromisoformat(value)+offset
                        else:shifted=date.fromisoformat(value)+offset
                        cell.value=shifted.isoformat() if isinstance(value,str) else shifted

        def add(ws,record):
            header=[cell.value for cell in ws[1]]
            if not set(record)<=set(header):raise ValueError(f'Unknown fields in {ws.title}')
            previous=ws.max_row
            ws.append([record.get(name,'') for name in header])
            for col in range(1,ws.max_column+1):
                source=ws.cell(previous,col);dest=ws.cell(ws.max_row,col)
                if source.has_style:dest._style=copy(source._style)
                dest.alignment=copy(source.alignment)
            for table in ws.tables.values():
                table.ref=f'A1:{ws.cell(ws.max_row,ws.max_column).coordinate}'

        def on(days):return (anchor+timedelta(days=days)).isoformat()
        disburse=books['Weekly_Disbursements.xlsx']
        add(disburse['Expected_Receipts'],dict(id='POC-GBP-RECEIPT',currency='GBP',amount=700000,value_date=on(13),
            description='SYNTHETIC confirmed sterling customer receipt, available before linked GBP payment',status='Confirmed'))
        add(disburse['Expected_Receipts'],dict(id='POC-CAD-RECEIPT',currency='CAD',amount=2000000,value_date=on(23),
            description='SYNTHETIC committed CAD inflow AFTER the earlier CAD cash trough',status='Confirmed'))
        for id,currency,amount,day,description in [
            ('POC-GBP-LINKED','GBP',700000,17,'SYNTHETIC GBP supplier payment linked to confirmed receipt and planned USD-funded purchase'),
            ('POC-GBP-LATER','GBP',3400000,26,'SYNTHETIC later GBP obligations; compare GBP funding sales with CAD'),
            ('POC-USD-NEED','USD',350000,27,'SYNTHETIC USD payment; compare funding currencies without duplicating same need'),
            ('POC-USD-FUTURE','USD',2000000,25,'SYNTHETIC future USD vendor payment with GBP budget; assess optional partial forward')]:
            add(payments,dict(id=id,currency=currency,amount=amount,value_date=on(day),category='Bill',description=description,status='Committed'))

        simulation=books['FX_Simulation.xlsx']
        add(simulation['Receipt_Timing'],dict(id='POC-TIMING-GBP',currency='GBP',synthetic=True,
            description='SYNTHETIC confirmed receipt timing; one-day delay still before payment',receipt_id='POC-GBP-RECEIPT',
            confirmation_reference='POC-CONF-GBP',delay_days=1,arrival_hour=11))
        orders=simulation['FX_Orders']
        # The existing planned GBP purchase now competes with the incoming GBP receipt.
        cols={cell.value:cell.column for cell in orders[1]}
        for row in range(2,orders.max_row+1):
            if orders.cell(row,cols['id']).value=='SIM-FX-PURCHASE-01':
                for field,value in dict(source_amount=969307.45,target_amount=700000,settlement_date=on(14),obligation_id='POC-GBP-LINKED',fee_jmd=76030.75).items():
                    orders.cell(row,cols[field]).value=value
                break
        else:raise ValueError('Missing original planned GBP purchase')
        for name in files:books[name].save(folder/name)
        return True
    finally:
        for wb in books.values():wb.close()


if __name__=='__main__':
    from .seed import SHARED
    print('Extended synthetic portfolio' if extend(SHARED) else 'Synthetic portfolio already extended')
