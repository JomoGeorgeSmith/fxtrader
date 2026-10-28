import copy
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from unittest.mock import patch
from backend.seed import seed, ROOT
from backend.app import read_context, digest
from backend.calculations import evaluate
from backend.liquidity_metrics import calculate_metrics

class TreasuryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.workspace = tempfile.TemporaryDirectory()
        cls.shared = Path(cls.workspace.name)
        with patch('backend.seed.SHARED',cls.shared):
            seed()
        cls.data = read_context(cls.shared)
        cls.rules = json.loads((ROOT/'policies/demo-rules.json').read_text(encoding='utf-8'))
        cls.now = datetime.fromisoformat(cls.data['accounts'][0]['as_of'])

    @classmethod
    def tearDownClass(cls):
        cls.workspace.cleanup()

    def test_metric_arithmetic_and_status(self):
        lcr,nsfr=calculate_metrics(self.data['liquidity_inputs'],self.rules)
        self.assertAlmostEqual(lcr['value_pct'],138.8888889)
        self.assertEqual(lcr['status'],'good')
        self.assertAlmostEqual(nsfr['value_pct'],111.1111111)
        self.assertEqual(nsfr['status'],'normal')

    def test_lcr_caps_inflows(self):
        rows=copy.deepcopy(self.data['liquidity_inputs'])
        next(r for r in rows if r['component']=='eligible_inflows_30d')['amount']=2000000000
        lcr,_=calculate_metrics(rows,self.rules)
        self.assertEqual(lcr['denominator'],450000000)

    def test_lcr_includes_maturity_mismatch_addon(self):
        rows=copy.deepcopy(self.data['liquidity_inputs'])
        next(r for r in rows if r['component']=='maturity_mismatch_addon')['amount']=100000000
        lcr,_=calculate_metrics(rows,self.rules)
        self.assertEqual(lcr['denominator'],1000000000)
        self.assertEqual(lcr['value_pct'],125)

    def test_metric_status_boundaries_and_findings(self):
        for amount,status in [(899000000,'at-risk'),(900000000,'normal'),(1080000000,'good')]:
            data=copy.deepcopy(self.data)
            next(r for r in data['liquidity_inputs'] if r['component']=='eligible_hqla')['amount']=amount
            result=evaluate(data,self.rules,self.now)
            self.assertEqual(result['metrics'][0]['status'],status)
            self.assertEqual(any(f['id']=='metric-lcr' for f in result['findings']),status=='at-risk')

    def test_zero_denominator_rejected(self):
        rows=copy.deepcopy(self.data['liquidity_inputs'])
        next(r for r in rows if r['component']=='required_stable_funding')['amount']=0
        with self.assertRaises(ValueError):
            calculate_metrics(rows,self.rules)

    def test_liquidity_sheet_migration_preserves_user_edits(self):
        from backend.seed import ensure_liquidity_inputs
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)
            path=target/'Bank_Balances.xlsx'
            shutil.copy2(self.shared/path.name,path)
            wb=load_workbook(path)
            wb['Accounts']['D3']=7777777
            del wb['Liquidity_Inputs']
            wb.save(path)
            wb.close()
            with patch('backend.seed.SHARED',target):
                ensure_liquidity_inputs()
            wb=load_workbook(path)
            self.assertEqual(wb['Accounts']['D3'].value,7777777)
            self.assertIn('Liquidity_Inputs',wb.sheetnames)
            wb.close()

    def test_seed_has_all_currencies(self):
        self.assertEqual({r['currency'] for r in self.data['accounts']},{'JMD','USD','GBP','CAD'})

    def test_forecast_and_reserve_are_separate(self):
        result=evaluate(self.data,self.rules,self.now)
        usd=next(p for p in result['positions'] if p['currency']=='USD')
        self.assertEqual(usd['available'],2300000)
        self.assertEqual(usd['projected_low'],-300000)
        self.assertEqual(usd['reserve'],460000)
        self.assertEqual(result['findings'][0]['id'],'liquidity-USD')

    def test_exposure_calculation(self):
        result=evaluate(self.data,self.rules,self.now)
        self.assertAlmostEqual(result['exposure_pct'],15.34963)

    def test_suggestion_cap_and_evidence(self):
        result=evaluate(self.data,self.rules,self.now)
        self.assertNotIn('suggestions',result, 'Calculation layer must not generate final cards')
        self.assertTrue(all(f['sources'] and f['rule'] for f in result['findings']))

    def test_funding_due_before_first_buffer_breach_not_weekly_low(self):
        from datetime import timedelta
        result=evaluate(self.data,self.rules,self.now)
        funding=next(f for f in result['findings'] if f['id']=='liquidity-USD')
        first=(self.now+timedelta(days=3)).date().isoformat()
        self.assertIn(f'by {first}',funding['summary'])

    def test_settled_payments_excluded(self):
        data=copy.deepcopy(self.data)
        for row in data['payments']:
            row['status']='Settled'
        result=evaluate(data,self.rules,self.now)
        self.assertFalse(any(f['id']=='liquidity-USD' for f in result['findings']))

    def test_three_week_request_is_visible_only_in_extended_forecast(self):
        from datetime import timedelta
        data=copy.deepcopy(self.data)
        data['payments']=[dict(id='FUTURE',currency='USD',amount=2000000,value_date=(self.now+timedelta(days=21)).date().isoformat(),status='Committed',_source='synthetic request')]
        data['receipts']=[]
        short=copy.deepcopy(self.rules)
        short['forecast_days']=7
        self.assertFalse(any(f['id']=='liquidity-USD' for f in evaluate(data,short,self.now)['findings']))
        long=copy.deepcopy(short)
        long['forecast_days']=28
        result=evaluate(data,long,self.now)
        self.assertTrue(any(f['id']=='liquidity-USD' for f in result['findings']))
        self.assertEqual(len([r for r in result['forecast'] if r['currency']=='USD']),29)

    def test_expected_deposits_are_not_counted_in_cautious_forecast(self):
        data=copy.deepcopy(self.data)
        data['payments']=[]
        data['receipts']=[dict(id='UNCERTAIN',currency='USD',amount=800000,value_date=self.now.date().isoformat(),status='Expected',_source='synthetic deposit')]
        result=evaluate(data,self.rules,self.now)
        day=next(r for r in result['forecast'] if r['currency']=='USD')
        self.assertEqual(day['balance'],3100000)
        self.assertEqual(day['confirmed_receipts_balance'],2300000)

    def test_scenario_addition_preserves_existing_rows_and_rejects_duplicate_apply(self):
        import shutil
        from backend.funding_scenario import apply_scenario
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)
            for file in self.shared.glob('*.xlsx'):
                shutil.copy2(file,target/file.name)
            before=read_context(target)
            manifest=apply_scenario(target,target/'runtime',self.now.date())
            after=read_context(target)
            self.assertTrue(Path(manifest['backup']).exists())
            for name in before:
                self.assertEqual(after[name][:len(before[name])],before[name])
            self.assertEqual(len(after['payments'])-len(before['payments']),9)
            self.assertEqual(len(after['receipts'])-len(before['receipts']),9)
            with self.assertRaises(ValueError):
                apply_scenario(target,target/'runtime',self.now.date())

    def test_semantic_digest_ignores_order_and_row_references(self):
        data=copy.deepcopy(self.data)
        data['accounts'].reverse()
        data['accounts'][0]['_source']='a different row'
        self.assertEqual(digest(data),digest(self.data))
        data['accounts'][0]['balance']+=1
        self.assertNotEqual(digest(data),digest(self.data))

    def test_formatting_save_does_not_change_digest(self):
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)
            for file in self.shared.glob('*.xlsx'):
                shutil.copy2(file,target/file.name)
            before=digest(read_context(target))
            wb=load_workbook(target/'Bank_Balances.xlsx')
            wb['Accounts']['D2'].font=Font(bold=True,color='FF0000')
            wb.save(target/'Bank_Balances.xlsx')
            wb.close()
            self.assertEqual(before,digest(read_context(target)))

    def test_invalid_currency_rejected(self):
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)
            for file in self.shared.glob('*.xlsx'):
                shutil.copy2(file,target/file.name)
            wb=load_workbook(target/'Bank_Balances.xlsx')
            wb['Accounts']['C2']='EUR'
            wb.save(target/'Bank_Balances.xlsx')
            wb.close()
            with self.assertRaises(ValueError):
                read_context(target)

if __name__=='__main__':
    unittest.main()
